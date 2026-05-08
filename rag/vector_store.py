"""ChromaDB-backed document indexing and retrieval for RAG."""

from __future__ import annotations

import hashlib
import os
import re
import unicodedata
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Iterable, Optional

from dotenv import load_dotenv

try:
    import chromadb
except ImportError:  # pragma: no cover
    chromadb = None

try:
    import fitz  # PyMuPDF
except ImportError:  # pragma: no cover
    fitz = None

try:
    from docx import Document as DocxDocument
except ImportError:  # pragma: no cover
    DocxDocument = None

try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
except ImportError:  # pragma: no cover
    RecursiveCharacterTextSplitter = None

try:
    from sentence_transformers import SentenceTransformer
except ImportError:  # pragma: no cover
    SentenceTransformer = None


load_dotenv()

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}
LETTER_RE = re.compile(r"[^\W\d_]{2,}", re.UNICODE)
PAGE_NUMBER_RE = re.compile(r"^(?:page\s*)?\d{1,4}(?:\s*(?:/|sur|of)\s*\d{1,4})?$", re.IGNORECASE)


def _resolve_project_path(value: Optional[Path | str], default: Path) -> Path:
    """Resolve relative env paths from the project root."""
    path = Path(value) if value else default
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    return path.resolve()


@dataclass(frozen=True)
class SearchResult:
    """One retrieved source chunk."""

    text: str
    metadata: dict
    score: float


@lru_cache(maxsize=2)
def _get_embedding_model(model_name: str) -> SentenceTransformer:
    """Load the embedding model once per process."""
    if SentenceTransformer is None:
        raise RuntimeError("sentence-transformers is not installed.")
    return SentenceTransformer(model_name)


def _sha256_bytes(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            h.update(block)
    return h.hexdigest()


def _sha256_text(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", text or "")
    normalized = re.sub(r"\s+", " ", normalized).strip().casefold()
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def clean_text(text: str) -> str:
    """Clean extracted text conservatively before chunking/embedding."""
    if not text:
        return ""

    text = unicodedata.normalize("NFKC", text)
    text = (
        text.replace("\u00a0", " ")
        .replace("\u200b", "")
        .replace("\ufeff", "")
        .replace("\u00ad", "")
    )

    chars: list[str] = []
    for ch in text:
        category = unicodedata.category(ch)
        if ch in "\n\t":
            chars.append(ch)
        elif category in {"Cc", "Co", "Cs"}:
            chars.append(" ")
        else:
            chars.append(ch)
    text = "".join(chars)

    text = re.sub(r"(?<=[^\W\d_])-\s*\n\s*(?=[^\W\d_])", "", text, flags=re.UNICODE)
    text = re.sub(r"[ \t]*\n[ \t]*", "\n", text)

    lines: list[str] = []
    for line in text.splitlines():
        line = re.sub(r"[ \t]{2,}", " ", line.strip())
        if not line:
            if lines and lines[-1] != "":
                lines.append("")
            continue
        if PAGE_NUMBER_RE.fullmatch(line):
            continue
        lines.append(line)

    return re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()


def is_informative(text: str, min_chars: int = 120, min_words: int = 8) -> bool:
    """Reject empty, tiny, numeric-only, or garbled text chunks."""
    value = (text or "").strip()
    if len(value) < min_chars:
        return False

    non_space = sum(1 for ch in value if not ch.isspace())
    letters = sum(1 for ch in value if ch.isalpha())
    digits = sum(1 for ch in value if ch.isdigit())
    private_use = sum(1 for ch in value if 0xE000 <= ord(ch) <= 0xF8FF)
    words = LETTER_RE.findall(value)

    if len(words) < min_words:
        return False
    if private_use / max(non_space, 1) > 0.02:
        return False
    if digits / max(non_space, 1) > 0.55:
        return False
    if letters / max(non_space, 1) < 0.20:
        return False

    return True


def _read_pdf(path: Path) -> list[dict]:
    if fitz is None:
        raise RuntimeError("PyMuPDF is not installed. Run: pip install PyMuPDF")

    pages: list[dict] = []
    with fitz.open(path) as doc:
        for page_number, page in enumerate(doc, start=1):
            text = page.get_text("text", sort=True) or ""
            text = clean_text(text)
            if is_informative(text, min_chars=40, min_words=3):
                pages.append({"text": text, "page": page_number})
    return pages


def _read_txt(path: Path) -> list[dict]:
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            text = path.read_text(encoding=encoding)
            text = clean_text(text)
            return [{"text": text, "page": -1}] if is_informative(text, min_chars=40, min_words=3) else []
        except UnicodeDecodeError:
            continue
    return []


def _read_docx(path: Path) -> list[dict]:
    if DocxDocument is None:
        raise RuntimeError("python-docx is not installed. Run: pip install python-docx")

    doc = DocxDocument(path)
    blocks = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    text = clean_text("\n".join(blocks))
    return [{"text": text, "page": -1}] if is_informative(text, min_chars=40, min_words=3) else []


def load_document(path: Path) -> list[dict]:
    """Load one supported document into page-like text records."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix in {".txt", ".md"}:
        return _read_txt(path)
    if suffix == ".docx":
        return _read_docx(path)
    return []


def discover_documents(data_dir: Path) -> list[Path]:
    """Find supported documents under data_dir."""
    if not data_dir.exists():
        return []
    return sorted(
        path
        for path in data_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


class ChromaRAGStore:
    """Small ChromaDB wrapper for indexing and semantic retrieval."""

    def __init__(
        self,
        data_dir: Optional[Path | str] = None,
        persist_dir: Optional[Path | str] = None,
        collection_name: Optional[str] = None,
        embedding_model_name: Optional[str] = None,
    ) -> None:
        if chromadb is None:
            raise RuntimeError("chromadb is not installed.")

        self.data_dir = _resolve_project_path(
            data_dir or os.getenv("RAG_DATA_DIR"),
            PROJECT_ROOT / "data-RAG",
        )
        self.persist_dir = _resolve_project_path(
            persist_dir or os.getenv("CHROMA_PERSIST_DIR"),
            PROJECT_ROOT / ".chroma",
        )
        self.collection_name = collection_name or os.getenv("CHROMA_COLLECTION", "agroscan_rag")
        self.embedding_model_name = embedding_model_name or os.getenv(
            "EMBEDDING_MODEL", "sentence-transformers/all-MiniLM-L6-v2"
        )

        chroma_host = os.getenv("CHROMA_HOST")
        if chroma_host:
            chroma_port = int(os.getenv("CHROMA_PORT", "8000"))
            self.client = chromadb.HttpClient(host=chroma_host, port=chroma_port)
        else:
            self.persist_dir.mkdir(parents=True, exist_ok=True)
            self.client = chromadb.PersistentClient(path=str(self.persist_dir))

        self.collection = self.client.get_or_create_collection(
            name=self.collection_name,
            metadata={"hnsw:space": "cosine"},
        )

        self.splitter = self._build_splitter()

    def _build_splitter(self):
        if RecursiveCharacterTextSplitter is None:
            return None
        return RecursiveCharacterTextSplitter(
            chunk_size=int(os.getenv("RAG_CHUNK_SIZE", "800")),
            chunk_overlap=int(os.getenv("RAG_CHUNK_OVERLAP", "120")),
            length_function=len,
            separators=["\n\n", "\n", ". ", "; ", ", ", " ", ""],
        )

    def _embed(self, texts: list[str]) -> list[list[float]]:
        model = _get_embedding_model(self.embedding_model_name)
        vectors = model.encode(
            texts,
            batch_size=int(os.getenv("EMBEDDING_BATCH_SIZE", "64")),
            show_progress_bar=False,
            convert_to_numpy=True,
            normalize_embeddings=True,
        )
        return vectors.tolist()

    def _split_text(self, text: str) -> list[str]:
        if self.splitter is not None:
            return self.splitter.split_text(text)

        size = int(os.getenv("RAG_CHUNK_SIZE", "800"))
        overlap = int(os.getenv("RAG_CHUNK_OVERLAP", "120"))
        chunks = []
        start = 0
        while start < len(text):
            chunks.append(text[start : start + size])
            start += max(size - overlap, 1)
        return chunks

    def count(self) -> int:
        """Return number of chunks in the collection."""
        return int(self.collection.count())

    def iter_document_chunks(self, path: Path) -> Iterable[dict]:
        """Yield validated chunks with source metadata for one document."""
        file_hash = _sha256_bytes(path)
        for page in load_document(path):
            for chunk_index, chunk in enumerate(self._split_text(page["text"])):
                chunk = clean_text(chunk)
                if not is_informative(chunk):
                    continue
                chunk_hash = _sha256_text(chunk)
                yield {
                    "id": f"{file_hash[:16]}-{chunk_index:05d}-{chunk_hash[:12]}",
                    "text": chunk,
                    "metadata": {
                        "filename": path.name,
                        "source_path": str(path.resolve()),
                        "file_hash": file_hash,
                        "chunk_hash": chunk_hash,
                        "chunk_index": chunk_index,
                        "page": int(page.get("page", -1)),
                    },
                }

    def index_documents(self, data_dir: Optional[Path | str] = None, batch_size: int = 64) -> dict:
        """Index all supported documents. Deterministic IDs make reruns idempotent."""
        root = Path(data_dir).resolve() if data_dir else self.data_dir
        documents = discover_documents(root)
        summary = {
            "data_dir": str(root),
            "files_seen": len(documents),
            "files_indexed": 0,
            "chunks_indexed": 0,
            "files_skipped": 0,
            "errors": [],
        }

        batch: list[dict] = []

        def flush() -> None:
            if not batch:
                return
            texts = [item["text"] for item in batch]
            self.collection.upsert(
                ids=[item["id"] for item in batch],
                documents=texts,
                embeddings=self._embed(texts),
                metadatas=[item["metadata"] for item in batch],
            )
            summary["chunks_indexed"] += len(batch)
            batch.clear()

        for path in documents:
            try:
                chunks = list(self.iter_document_chunks(path))
                if not chunks:
                    summary["files_skipped"] += 1
                    continue
                summary["files_indexed"] += 1
                for chunk in chunks:
                    batch.append(chunk)
                    if len(batch) >= batch_size:
                        flush()
            except Exception as exc:  # keep indexing resilient across bad PDFs
                summary["errors"].append({"file": str(path), "error": str(exc)})

        flush()
        return summary

    def search(self, query: str, top_k: int = 4) -> list[SearchResult]:
        """Retrieve semantically relevant chunks for query."""
        query = (query or "").strip()
        if not query or self.count() == 0:
            return []

        raw = self.collection.query(
            query_embeddings=self._embed([query]),
            n_results=min(top_k, self.count()),
            include=["documents", "metadatas", "distances"],
        )

        documents = (raw.get("documents") or [[]])[0]
        metadatas = (raw.get("metadatas") or [[]])[0]
        distances = (raw.get("distances") or [[]])[0]

        results: list[SearchResult] = []
        for text, metadata, distance in zip(documents, metadatas, distances):
            results.append(
                SearchResult(
                    text=text,
                    metadata=metadata or {},
                    score=max(0.0, 1.0 - float(distance)),
                )
            )
        return results
